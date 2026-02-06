import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import threading
from concurrent.futures import ThreadPoolExecutor
import pdfplumber
import numpy as np
import pandas as pd
from pdf2image import convert_from_path
from PIL import Image
import re
import torch

# Surya OCR
from surya.foundation import FoundationPredictor
from surya.detection import DetectionPredictor
from surya.recognition import RecognitionPredictor


# ---------------- CONFIG ----------------
DPI = 600
LINE_TOL = 2.0

MIN_LINE_LEN = 50


# ---------------- HELPERS ----------------
def clean_extracted_text(text):
    """
    ✅ Enhanced Kaggle text cleaning implementation
    Applies all regex fixes from the working Kaggle notebook
    """
    if not text:
        return ""
    
    # 1. REMOVE NOISE & TAGS
    text = re.sub(r"<[^>]+>", " ", text)  # HTML tags
    text = re.sub(r"\bPhoto\b|\bAvailable\b", " ", text)  # Photo/Available text
    text = re.sub(r"[·•]", " ", text)  # Bullet points and dots
    
    # 2. FIX SPELLING: Change "लिग :", "लीग :", "लंग :" to "लिंग :"
    text = re.sub(r"(?:लिग|लीग|लंग)\s*:", "लिंग :", text)
    
    # 2b. FIX AGE SPELLING: Change "विय :" to "वय :"
    text = re.sub(r"विय\s*:", "वय :", text)
    
    # 3. GENDER FIX (Now works because spelling is fixed)
    def fix_gender(match):
        val = match.group(1)
        return "लिंग : स्त्री" if "स्त्री" in val else "लिंग : पु"
    
    text = re.sub(r"लिंग\s*:\s*([^\s|,\n]+)", fix_gender, text)
    
    # 4. AGE NUMBER FIX (8->४, 0->७, etc)
    def fix_age_digits(match):
        prefix = match.group(1)
        digits = match.group(2)
        mapping = {'8': '४', '9': '९', '0': '७', '4': '५', '3': '३', '2': '२'}
        new_digits = "".join([mapping.get(c, c) for c in digits])
        return prefix + new_digits
    
    text = re.sub(r"(वय\s*:\s*)([\d]+)", fix_age_digits, text)
    
    # 5. GLOBAL QUESTION MARK FIX (? -> २)
    text = re.sub(r"[\d२-९]*\?+[\d२-९]*", lambda m: m.group(0).replace("?", "२"), text)
    
    # 6. NORMALIZE WHITESPACE
    text = re.sub(r"\s+", " ", text)
    
    return text


# ---------------- PARSING (KAGGLE-FAITHFUL) ----------------
def parse_header(text):
    """Parse header text into structured data - only extracts fields that exist"""
    data = {}
    
    # --- DIVISION (निवडणूक विभाग) ---
    # Pattern 1: "विभाग : 5" (File 1 format)
    div_match = re.search(r"निवडणूक\s+विभाग\s*:\s*([^निवार्चन]+?)(?=\s*निवार्चन|$)", text)
    if not div_match:
        # Pattern 2: "प्रभाग क्र: 5" (File 2 format)
        div_match = re.search(r"प्रभाग\s+क्र\s*[:\s]+(\d+)", text)
    data['division'] = div_match.group(1).strip() if div_match else ""
    
    # --- ELECTORAL CONSTITUENCY (निवार्चन गण) ---
    # ONLY extract if "निवार्चन गण" explicitly exists
    # Pattern: "निवार्चन गण: 9"
    gan_match = re.search(r"निवार्चन\s+गण\s*:\s*(\d+)", text)
    data['gan'] = gan_match.group(1).strip() if gan_match else ""
    
    # --- PART NUMBER (यादी भाग क्र.) ---
    # Captures entire text from "यादी भाग क्र." until "पत्ता" or end of header
    # Example: "यादी भाग क्र. १४३ : १ - राहुल नगर स्वातंत्र्त्र सैनिक कॉलोनी परभणी शहर"
    part_match = re.search(r"(यादी\s+भाग\s+क्र\..*?)(?=\s*पत्ता|$)", text)
    data['part_no'] = part_match.group(1).strip() if part_match else ""
    
    # --- ADDRESS (पत्ता) ---
    # Only extracts if "पत्ता :" explicitly exists
    # Pattern: "पत्ता : जि.प.प्रा.शाळा घेवंडा"
    addr_match = re.search(r"पत्ता\s*:\s*(.+?)(?=\s*मतदान|$)", text)
    data['address'] = addr_match.group(1).strip() if addr_match else ""
    
    # --- POLLING STATION (मतदान केंद्र) ---
    # Pattern: "मतदान केंद्र : 9 Ghevanda"
    poll_match = re.search(r"मतदान\s+केंद्र\s*:\s*(.+?)(?=\s*$)", text)
    data['polling_station'] = poll_match.group(1).strip() if poll_match else ""
    
    return data


def parse_box_text(text, header_data):
    """Parse box text into structured voter data with robust field extraction"""
    row_data = {}
    
    # Define field keywords for boundary detection
    FIELD_KEYWORDS = [
        'मतदाराचे पूर्ण', 'घर क्रमांक', 'लिंग', 'वय',
        'नाव', 'नांव',  # Name keywords that appear after some fields
        'voter', 'house', 'gender', 'age'
    ]
    
    def extract_until_delimiter(pattern, text_to_search):
        """
        Extract text after pattern until we hit:
        1. Pipe delimiter |
        2. Next field keyword
        3. End of line
        """
        match = re.search(pattern, text_to_search, re.IGNORECASE)
        if not match:
            return ""
        
        # Get text after the matched pattern
        start_pos = match.end()
        remaining_text = text_to_search[start_pos:]
        
        # Find the earliest delimiter
        delimiters = []
        
        # Check for pipe
        pipe_pos = remaining_text.find('|')
        if pipe_pos != -1:
            delimiters.append(pipe_pos)
        
        # Check for next field keyword
        for keyword in FIELD_KEYWORDS:
            keyword_pos = remaining_text.lower().find(keyword.lower())
            if keyword_pos != -1 and keyword_pos > 0:  # Ignore if at start
                delimiters.append(keyword_pos)
        
        # Use earliest delimiter, or take all remaining text
        if delimiters:
            end_pos = min(delimiters)
            extracted = remaining_text[:end_pos]
        else:
            extracted = remaining_text
        
        return extracted.strip()
    
    # --- VOTER ID & S-NUMBER EXTRACTION (Robust - handles any position) ---
    # Try to find voter ID (alphanumeric like WMJ6725378 or NMG6681910)
    voter_id_match = re.search(r"\b([A-Z]{2,}[A-Z0-9]{5,})\b", text)
    row_data['voter_id'] = voter_id_match.group(1) if voter_id_match else ""
    
    # S-number (pattern like 95/153/1)
    s_match = re.search(r"(\d+/\d+/\d+)", text)
    row_data['s'] = s_match.group(1) if s_match else ""
    
    # --- SERIAL NUMBER EXTRACTION ---
    # Try pipe-separated first: | 123 |
    sr_match = re.search(r"\|\s*(\d+)\s*\|", text)
    if sr_match:
        row_data['sr.no'] = sr_match.group(1)
    else:
        # Fallback: use last part of s-number (e.g., 95/153/1 -> 1)
        if row_data['s']:
            parts = row_data['s'].split('/')
            if len(parts) == 3:
                row_data['sr.no'] = parts[2]
            else:
                row_data['sr.no'] = ""
        else:
            row_data['sr.no'] = ""
    
    # --- HEADER COLUMNS (only include fields that exist) ---
    row_data['निवार्चन गण'] = header_data.get('gan', '')
    row_data['यादी भाग क्र.'] = header_data.get('part_no', '')
    row_data['पत्ता'] = header_data.get('address', '')
    row_data['मतदान केंद्र'] = header_data.get('polling_station', '')
    
    # --- NAME EXTRACTION (Robust) ---
    row_data['मतदाराचे पूर्ण'] = extract_until_delimiter(
        r"मतदाराचे\s+पूर्ण[:\s]*",
        text
    )
    
    # --- HOUSE NUMBER EXTRACTION (Robust) ---
    row_data['घर क्रमांक'] = extract_until_delimiter(
        r"घर\s+क्रमां?क\s*[:.\s]*",
        text
    )
    
    # --- GENDER EXTRACTION (Robust) ---
    row_data['लिंग'] = extract_until_delimiter(
        r"लिंग\s*[:.\s]*",
        text
    )
    
    # --- AGE EXTRACTION (Robust) ---
    age_text = extract_until_delimiter(
        r"वय\s*[:.\s]*",
        text
    )
    # Extract only digits (handles both English and Devanagari numerals)
    age_match = re.search(r"([\d०-९]+)", age_text)
    row_data['वय'] = age_match.group(1) if age_match else age_text
    
    # --- RAW HEADER (Full header text for reference) ---
    row_data['header'] = header_data.get('raw_header', '')
    
    # --- DELETED FLAG ---
    row_data['is_deleted'] = "**" if "**" in text else ""
    
    return row_data


def get_boxes_and_header(pdf_path, page_num, img_w, img_h):
    """
    ✅ VERIFIED Kaggle implementation
    Uses proper pdfplumber keys: 'top', 'bottom', 'x0', 'x1'
    """
    with pdfplumber.open(pdf_path) as pdf:
        page = pdf.pages[page_num]
        lines = page.lines

        def get_filtered_coords(lines_subset, coord_key, len_key, tol=2.0, min_len=50):
            """Cluster lines and filter by max length in cluster"""
            lines_subset.sort(key=lambda x: x[coord_key])
            if not lines_subset:
                return []

            unique_coords = []
            curr_pos = [lines_subset[0][coord_key]]
            curr_len = [abs(lines_subset[0][len_key[1]] - lines_subset[0][len_key[0]])]

            for l in lines_subset[1:]:
                pos = l[coord_key]
                length = abs(l[len_key[1]] - l[len_key[0]])

                if abs(pos - curr_pos[-1]) <= tol:
                    curr_pos.append(pos)
                    curr_len.append(length)
                else:
                    if max(curr_len) > min_len:
                        unique_coords.append(sum(curr_pos) / len(curr_pos))
                    curr_pos = [pos]
                    curr_len = [length]

            if max(curr_len) > min_len:
                unique_coords.append(sum(curr_pos) / len(curr_pos))

            return unique_coords

        # ✅ CORRECT: Use 'top'/'bottom' and 'x0'/'x1' (guaranteed by pdfplumber)
        v_raw = [l for l in lines if abs(l["x1"] - l["x0"]) < 1]
        h_raw = [l for l in lines if abs(l["bottom"] - l["top"]) < 1]

        unique_vs = get_filtered_coords(v_raw, "x0", ("top", "bottom"))
        unique_hs = get_filtered_coords(h_raw, "top", ("x0", "x1"))

        if not unique_vs or not unique_hs:
            return [], None

        scale_x = img_w / page.width
        scale_y = img_h / page.height

        # Header
        header_bottom = int(unique_hs[0] * scale_y)
        header_rect = (0, 0, img_w, max(50, header_bottom - 5))

        boxes = []
        for i in range(len(unique_hs) - 1):
            y_top = unique_hs[i]
            y_bottom = unique_hs[i + 1]
            if y_bottom - y_top < 40:
                continue

            for j in range(len(unique_vs) - 1):
                x_left = unique_vs[j]
                x_right = unique_vs[j + 1]
                if x_right - x_left < 100:
                    continue

                x = int(x_left * scale_x)
                y = int(y_top * scale_y)
                w = int((x_right - x_left) * scale_x)
                h = int((y_bottom - y_top) * scale_y)

                boxes.append((x, y, w, h))

        boxes.sort(key=lambda b: (b[1] // 50, b[0]))
        return boxes, header_rect


# ---------------- GUI APP ----------------
class VoterExtractorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Voter PDF OCR Extractor (Batch Mode)")
        # ✅ Increased window size for 3x bigger UI
        self.root.geometry("2000x1600")

        # ✅ Configure Large Styles (3x bigger)
        style = ttk.Style()
        style.configure("Big.TLabel", font=("Helvetica", 20))
        style.configure("Big.TButton", font=("Helvetica", 24))
        style.configure("Big.TEntry", font=("Helvetica", 24))
        style.configure("Big.TCheckbutton", font=("Helvetica", 18))
        style.configure("Big.TLabelframe.Label", font=("Helvetica", 24, "bold"))

        # Batch Processing Variables
        self.file_paths = []  # List of selected files
        self.threads_var = tk.StringVar(value="4")
        
        # Device Selection
        self.device_var = tk.StringVar(value="auto")
        self.available_devices = self.detect_devices()
        
        # Column Selection (all 13 columns available)
        self.all_columns = [
            ('sr.no', 'Serial Number'),
            ('मतदाराचे पूर्ण', 'Full Name'),
            ('लिंग', 'Gender'),
            ('वय', 'Age'),
            ('s', 'S-Number'),
            ('voter_id', 'Voter ID'),
            ('निवार्चन गण', 'Electoral Constituency'),
            ('मतदान केंद्र', 'Polling Station'),
            ('पत्ता', 'Address'),
            ('यादी भाग क्र.', 'Part Number'),
            ('घर क्रमांक', 'House Number'),
            ('header', 'Raw Header'),
            ('is_deleted', 'Deleted Flag')
        ]
        
        # Column checkbox variables (all enabled by default)
        self.column_vars = {}
        for col_id, col_name in self.all_columns:
            self.column_vars[col_id] = tk.BooleanVar(value=True)
        
        # Build UI first (creates log_area)
        self.build_ui()
        
        # Load Surya OCR models with selected device
        self.load_models()

        self.ocr_lock = threading.Lock()
    
    # ---------- DEVICE DETECTION ----------
    def detect_devices(self):
        """Detect available compute devices (CPU, CUDA, MPS)"""
        devices = []
        
        # Always have CPU
        devices.append(("cpu", "CPU"))
        
        # Check for CUDA (NVIDIA GPU)
        if torch.cuda.is_available():
            gpu_count = torch.cuda.device_count()
            for i in range(gpu_count):
                gpu_name = torch.cuda.get_device_name(i)
                gpu_mem = torch.cuda.get_device_properties(i).total_memory / (1024**3)  # GB
                devices.append((f"cuda:{i}", f"GPU {i}: {gpu_name} ({gpu_mem:.1f}GB)"))
        
        # Check for MPS (Apple Silicon)
        if hasattr(torch.backends, 'mps') and torch.backends.mps.is_available():
            devices.append(("mps", "Apple Silicon GPU (MPS)"))
        
        return devices
    
    def load_models(self):
        """Load Surya OCR models on selected device"""
        device = self.device_var.get()
        
        # Extract device name if it's from combo box (format: "device - description")
        if " - " in device:
            device = device.split(" - ")[0]
        
        if device == "auto":
            # Auto-select: prefer CUDA > MPS > CPU
            if torch.cuda.is_available():
                device = "cuda:0"
            elif hasattr(torch.backends, 'mps') and torch.backends.mps.is_available():
                device = "mps"
            else:
                device = "cpu"
        
        # Safe logging (works even if UI not ready)
        def safe_log(msg):
            if hasattr(self, 'log_area'):
                self.log(msg)
            else:
                print(msg)
        
        safe_log(f"⏳ Loading Surya OCR models on {device.upper()}...")
        
        try:
            # Load models with explicit device
            self.foundation_predictor = FoundationPredictor(device=device)
            self.det_predictor = DetectionPredictor(device=device)
            # RecognitionPredictor doesn't accept device parameter, uses foundation_predictor's device
            self.rec_predictor = RecognitionPredictor(
                foundation_predictor=self.foundation_predictor
            )
            safe_log(f"✅ Models loaded successfully on {device.upper()}!")
        except Exception as e:
            safe_log(f"❌ Error loading models on {device}: {e}")
            safe_log("⚠️ Falling back to CPU...")
            self.foundation_predictor = FoundationPredictor(device="cpu")
            self.det_predictor = DetectionPredictor(device="cpu")
            self.rec_predictor = RecognitionPredictor(
                foundation_predictor=self.foundation_predictor
            )
            safe_log("✅ Models loaded on CPU")

    # ---------- UI ----------
    def build_ui(self):
        # Step 1: PDF Selection (List Mode)
        frame_top = ttk.LabelFrame(self.root, text="Step 1: Select PDFs (Batch Mode)", style="Big.TLabelframe")
        frame_top.pack(fill="x", padx=20, pady=20)

        # File List Button
        btn_frame = ttk.Frame(frame_top)
        btn_frame.pack(fill="x", padx=10, pady=5)
        
        ttk.Button(btn_frame, text="📂 Add PDF Files", command=self.add_files, style="Big.TButton").pack(side="left", padx=10)
        ttk.Button(btn_frame, text="📁 Add Folder", command=self.add_folder, style="Big.TButton").pack(side="left", padx=10)
        ttk.Button(btn_frame, text="🗑️ Clear List", command=self.clear_files, style="Big.TButton").pack(side="left", padx=10)
        
        # Listbox to show selected files
        self.file_listbox = tk.Listbox(frame_top, height=5, font=("Helvetica", 16))
        self.file_listbox.pack(fill="x", padx=10, pady=10)

        # Step 2: Settings
        frame_settings = ttk.LabelFrame(self.root, text="Step 2: Settings", style="Big.TLabelframe")
        frame_settings.pack(fill="x", padx=20, pady=20)
        
        # DPI info
        ttk.Label(frame_settings, text=f"DPI: {DPI}", style="Big.TLabel").pack(side="left", padx=20, pady=10)
        
        # Threads
        ttk.Label(frame_settings, text="Threads:", style="Big.TLabel").pack(side="left", padx=20)
        ttk.Spinbox(frame_settings, from_=1, to=16, textvariable=self.threads_var, width=5, font=("Helvetica", 24)).pack(side="left", padx=10)
        
        # Device Selection
        ttk.Label(frame_settings, text="Device:", style="Big.TLabel").pack(side="left", padx=20)
        device_options = [("auto", "Auto (Recommended)")] + self.available_devices
        device_combo = ttk.Combobox(
            frame_settings,
            textvariable=self.device_var,
            values=[f"{name} - {desc}" for name, desc in device_options],
            state="readonly",
            width=30,
            font=("Helvetica", 18)
        )
        device_combo.pack(side="left", padx=10)
        device_combo.current(0)  # Default to Auto
        


        
        # Apply to all checkbox (Still useful for future settings?)
        # Keeping it for now as it doesn't hurt, or we can remove if 'Start/End' was its only use.
        # Actually, let's keep it minimal as per user request for "same as code".
        # But wait, Apply to All was for SKIP settings. If SKIP is gone, this is useless.
        # Removing Apply to All as well.

        # Step 3: Column Selector
        frame_columns = ttk.LabelFrame(self.root, text="Step 3: Select Columns to Export", style="Big.TLabelframe")
        frame_columns.pack(fill="x", padx=20, pady=20)
        
        ttk.Label(frame_columns, text="✅ Choose which columns to include in Excel:", style="Big.TLabel").pack(padx=20, pady=10)
        
        # Create grid for checkboxes (3 columns)
        checkbox_frame = ttk.Frame(frame_columns)
        checkbox_frame.pack(padx=20, pady=10)
        
        for idx, (col_id, col_name) in enumerate(self.all_columns):
            row = idx // 3
            col = idx % 3
            cb = ttk.Checkbutton(
                checkbox_frame,
                text=f"{col_id} ({col_name})",
                variable=self.column_vars[col_id],
                style="Big.TCheckbutton"
            )
            cb.grid(row=row, column=col, sticky="w", padx=10, pady=5)
        
        # Select All / Deselect All buttons
        button_frame = ttk.Frame(frame_columns)
        button_frame.pack(pady=10)
        
        def select_all():
            for var in self.column_vars.values():
                var.set(True)
        
        def deselect_all():
            for var in self.column_vars.values():
                var.set(False)
        
        ttk.Button(button_frame, text="Select All", command=select_all).pack(side="left", padx=5)
        ttk.Button(button_frame, text="Deselect All", command=deselect_all).pack(side="left", padx=5)

        # Step 4: Start Button (Moved down)
        ttk.Button(
            self.root, text="🚀 Start Extraction", command=self.start_processing, style="Big.TButton"
        ).pack(pady=20)

        # Log Area
        self.log_area = scrolledtext.ScrolledText(
            self.root, height=18, font=("Consolas", 16)  # Larger font for log
        )
        self.log_area.pack(fill="both", expand=True, padx=20, pady=20)

    # ---------- LOG ----------
    def log(self, msg):
        self.log_area.insert(tk.END, msg + "\n")
        self.log_area.see(tk.END)
        self.root.update_idletasks()

    # ---------- FILE ----------
    def add_files(self):
        files = filedialog.askopenfilenames(filetypes=[("PDF Files", "*.pdf")])
        for f in files:
            if f not in self.file_paths:
                self.file_paths.append(f)
                self.file_listbox.insert(tk.END, f) # Simply show full path for clarity
    
    def add_folder(self):
        """Select a folder and add all PDF files in it"""
        import glob
        folder = filedialog.askdirectory(title="Select Folder Containing PDFs")
        if folder:
            # Find all PDF files in the folder
            pdf_files = glob.glob(os.path.join(folder, "*.pdf"))
            if not pdf_files:
                messagebox.showinfo("No PDFs Found", f"No PDF files found in:\n{folder}")
                return
            
            # Add all found PDFs
            added_count = 0
            for f in pdf_files:
                if f not in self.file_paths:
                    self.file_paths.append(f)
                    self.file_listbox.insert(tk.END, f)
                    added_count += 1
            
            messagebox.showinfo("Folder Added", f"Added {added_count} PDF file(s) from:\n{folder}")
    
    def clear_files(self):
        self.file_paths = []
        self.file_listbox.delete(0, tk.END)


    # ---------- PROCESS ----------
    def start_processing(self):
        if not self.file_paths:
            messagebox.showerror("Error", "Select at least one PDF first")
            return

        threading.Thread(target=self.process_batch, daemon=True).start()

    def process_batch(self):
        """
        NEW STRATEGY: Batch Processing
        - Iterate over all selected PDFs
        """
        total_files = len(self.file_paths)
        
        for file_idx, pdf_path in enumerate(self.file_paths):
            self.log(f"\n🚀 Processing File {file_idx + 1}/{total_files}: {pdf_path.split('/')[-1]}")
            
            # Reset results for this file (Sr.No starts at 1)
            all_results = []
            with pdfplumber.open(pdf_path) as p:
                total_pages = len(p.pages)
            
            self.log(f"ℹ️ Total Pages: {total_pages}. Scanning for voter data...")
            
            processing_started = False
            max_workers = int(self.threads_var.get())

            # Single Pass Loop
            for page_num in range(total_pages):
                self.log(f"   -> Checking Page {page_num + 1}...")
                
                # 1. Convert Image (High Res)
                try:
                    img = convert_from_path(pdf_path, dpi=DPI, first_page=page_num + 1, last_page=page_num + 1)[0]
                except Exception as e:
                    self.log(f"      ⚠️ Error converting page: {e}")
                    continue
                
                # 2. Detect Boxes
                boxes, header_rect = get_boxes_and_header(pdf_path, page_num, img.width, img.height)
                
                # 3. Validate Page
                is_valid_page = False
                if boxes:
                    try:
                        b = boxes[0]
                        crop = img.crop((b[0], b[1], b[0] + b[2], b[1] + b[3]))
                        with self.ocr_lock:
                            preds = self.rec_predictor([crop], det_predictor=self.det_predictor)
                        txt_chk = " ".join([l.text for l in preds[0].text_lines])
                        if re.search(r"[A-Z]{2,}\d+|[A-Z]+\d{3,}|[A-Z]+/[0-9]+/", txt_chk):
                            is_valid_page = True
                    except:
                        is_valid_page = False

                # 4. On-the-Go Logic
                if not processing_started:
                    if not is_valid_page:
                        self.log("      ⚠️ Intro/Non-Voter Page. Skipping.")
                        continue
                    else:
                        self.log("      ✅ START FOUND! Valid Voter ID Detected.")
                        processing_started = True
                else:
                    if not is_valid_page:
                        self.log("      🛑 END DETECTED. No valid boxes found. Stopping file.")
                        break

                # =========================================================
                # 5. PROCESS PAGE
                # =========================================================
                # =========================================================
                # 5. PROCESS PAGE
                # =========================================================
                self.log(f"   Processing Page {page_num + 1}...")
                
                # Note: 'boxes' and 'header_rect' are already computed in step 2
                
                if not boxes:
                    self.log(f"   ⚠️ Page {page_num + 1}: No boxes detected")
                    continue
                    self.log(f"   ⚠️ Page {page_num + 1}: No boxes detected")
                    continue
                
                # Step 3: Extract header
                header_data = {}
                if header_rect:
                    hc = img.crop(header_rect)
                    h_preds = self.rec_predictor([hc], det_predictor=self.det_predictor)
                    header_text = " ".join([l.text for l in h_preds[0].text_lines])
                    header_text = clean_extracted_text(header_text)
                    header_data['raw_header'] = header_text
                    parsed_header = parse_header(header_text)
                    header_data.update(parsed_header)
                
                # Step 4: Prepare box crops
                box_crops = []
                for b in boxes:
                    c = img.crop((b[0], b[1], b[0] + b[2], b[1] + b[3]))
                    if np.mean(np.array(c.convert("L"))) < 250:
                        box_crops.append(c)
                
                # Step 5: Process boxes in parallel
                with ThreadPoolExecutor(max_workers=max_workers) as executor:
                    futures = []
                    for i, crop in enumerate(box_crops):
                        future = executor.submit(self.process_single_box, crop, header_data, page_num, i)
                        futures.append(future)
                    
                    page_results = []
                    for future in futures:
                        row = future.result()
                        # Only add rows that have a voter ID (filter out headers/noise)
                        if row and row.get('voter_id', '').strip():
                            page_results.append(row)
                
                # Sort and Interpolate
                page_results.sort(key=lambda x: x['box_idx'])
                self.interpolate_serial_numbers(page_results)
                
                # Cleanup internal key
                for r in page_results:
                    r.pop('box_idx', None)
                    
                all_results.extend(page_results)
                self.log(f"   ✅ Page {page_num + 1} Done. Voters found: {len(page_results)}")
            # ✅ FINAL LOGIC: Overwrite Sr.No with sequential count
            for i, row in enumerate(all_results, 1):
                row['sr.no'] = str(i)
            
            
            # ✅ Save Excel with specific filename
            self.save_excel(all_results, source_pdf_path=pdf_path)
        
        self.log(f"\n🎉 BATCH COMPLETED! All PDF files processed.")
        messagebox.showinfo("Batch Complete", "All files have been processed successfully!")



    def process_single_box(self, crop, header_data, page_num, box_num):
        """
        Process a single box with OCR
        This runs in parallel across multiple threads
        """
        try:
            # Get thread ID for logging
            thread_id = threading.current_thread().name
            self.log(f"   🧵 Thread {thread_id}: Processing box {box_num + 1} on page {page_num + 1}")
            
            # OCR this single box (thread-safe with lock)
            with self.ocr_lock:
                preds = self.rec_predictor([crop], det_predictor=self.det_predictor)
            
            # Extract text - join with | to preserve field separators from PDF
            raw_txt = " | ".join([l.text for l in preds[0].text_lines])
            
            # ✅ Apply Kaggle-faithful text cleaning
            cleaned_txt = clean_extracted_text(raw_txt)
            
            # ✅ Parse into structured data (Kaggle-faithful)
            row = parse_box_text(cleaned_txt, header_data)
            
            # Store raw text for debugging/verification
            row['raw_text'] = raw_txt
            row['cleaned_text'] = cleaned_txt
            
            # Add box index for sorting
            row['box_idx'] = box_num
            
            self.log(f"   ✅ Thread {thread_id}: Completed box {box_num + 1}")
            return row
        except Exception as e:
            self.log(f"❌ Error processing box {box_num} on page {page_num + 1}: {e}")
            return None

    def interpolate_serial_numbers(self, results):
        """Fill in missing serial numbers by interpolating between known values"""
        if not results:
            return
        
        # Find indices with valid serial numbers
        valid_indices = []
        for i, row in enumerate(results):
            sr = row.get('sr.no', '')
            if sr and sr.isdigit():
                valid_indices.append((i, int(sr)))
        
        if len(valid_indices) < 2:
            return  # Need at least 2 points to interpolate
        
        # Interpolate between each pair of valid indices
        for j in range(len(valid_indices) - 1):
            start_idx, start_val = valid_indices[j]
            end_idx, end_val = valid_indices[j + 1]
            
            # Calculate step size
            gap = end_idx - start_idx
            if gap <= 1:
                continue  # No gap to fill
            
            value_diff = end_val - start_val
            step = value_diff / gap
            
            # Fill in missing values
            for k in range(1, gap):
                interpolated_val = int(start_val + (step * k))
                results[start_idx + k]['sr.no'] = str(interpolated_val)


    # ---------- EXCEL ----------
    # ---------- EXCEL ----------
    def save_excel(self, data, source_pdf_path=None):
        """Save extracted data to Excel with 13-column structure (Excel only)"""
        if not data:
            self.log("⚠️  No data to save")
            return

        self.log("📊 Formatting Excel...")
        
        # Get selected columns from checkboxes (order matches self.all_columns)
        all_possible_columns = [
            'sr.no',
            'मतदाराचे पूर्ण',
            'लिंग',
            'वय',
            's',
            'voter_id',
            'निवार्चन गण',
            'मतदान केंद्र',
            'पत्ता',
            'यादी भाग क्र.',
            'घर क्रमांक',
            'header',
            'is_deleted'
        ]
        
        # Filter to only selected columns
        final_columns = [col for col in all_possible_columns if self.column_vars[col].get()]
        
        # Log selected columns for debugging
        self.log(f"📋 Selected columns ({len(final_columns)}/{len(all_possible_columns)}): {', '.join(final_columns)}")
        
        if not final_columns:
            self.log("⚠️  No columns selected! Please select at least one column.")
            messagebox.showwarning("No Columns", "Please select at least one column to export.")
            return
        
        # Ensure all columns exist in data
        for row in data:
            for col in final_columns:
                if col not in row:
                    row[col] = ""
        
        # Generate output filename
        if source_pdf_path:
            output_file = source_pdf_path.replace(".pdf", "_Voter_List_Final.xlsx")
        else:
            output_file = "Voter_List_Final.xlsx"
        
        # Save using manual openpyxl (avoids sheet visibility error)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Voters"
            
            # Write header row (bold)
            for col_idx, col_name in enumerate(final_columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
            
            # Write data rows
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, col_name in enumerate(final_columns, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ''))
            
            # Auto-adjust column widths
            for col_idx, col_name in enumerate(final_columns, 1):
                max_len = len(col_name)
                for row_data in data:
                    val_len = len(str(row_data.get(col_name, '')))
                    if val_len > max_len:
                        max_len = val_len
                    ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 50)
            
            wb.save(output_file)
            self.log(f"✅ Excel saved: {output_file}")
            self.log(f"📊 Total voters extracted: {len(data)}")
            messagebox.showinfo("Success", f"✅ Saved {len(data)} voters to:\n{output_file}")
        except Exception as e:
            self.log(f"❌ Excel export failed: {e}")
            messagebox.showerror("Error", f"Failed to save Excel file:\n{e}")
    
    def save_raw_text(self, data, source_pdf_path=None):
        """Save raw OCR text in format: PAGE → HEADER → BOX X-Y"""
        if not data:
            return
        
        # Generate output filename
        if source_pdf_path:
            import os
            base_name = os.path.splitext(os.path.basename(source_pdf_path))[0]
            output_file = f"{base_name}_raw_text.txt"
        else:
            output_file = "raw_text_output.txt"
        
        try:
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write("--- RAW EXTRACTED VOTER DATA ---\n\n\n")
                
                # Group by page (using header as page indicator)
                current_page = 1
                current_header = None
                box_in_page = 1
                
                for row in data:
                    header = row.get('header', '')
                    
                    # New page detected (header changed)
                    if header != current_header:
                        if current_header is not None:
                            current_page += 1
                        current_header = header
                        box_in_page = 1
                        
                        # Write page header
                        f.write(f"=== PAGE {current_page} ===\n")
                        f.write(f"HEADER: {header if header else '(No header)'}\n\n")
                    
                    # Write box text - show BOTH raw and cleaned for comparison
                    f.write(f"BOX {box_in_page} (RAW): {row.get('raw_text', 'N/A')}\n")
                    f.write(f"BOX {box_in_page} (CLEANED - given to parser): {row.get('cleaned_text', 'N/A')}\n\n")
                    
                    box_in_page += 1
            
            self.log(f"📄 Raw text saved: {output_file}")
        except Exception as e:
            self.log(f"❌ Error saving raw text: {e}")


# ---------------- RUN ----------------
if __name__ == "__main__":
    root = tk.Tk()
    app = VoterExtractorApp(root)
    root.mainloop()
