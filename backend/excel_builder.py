"""
Excel Builder Module — Generates formatted .xlsx files from parsed voter data.
"""

import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .data_parser import FINAL_COLUMNS

logger = logging.getLogger(__name__)


def build_excel(all_results: list[dict], output_path: str) -> str:
    """
    Create a formatted Excel workbook from parsed voter records.

    Args:
        all_results: List of voter record dicts (from parse_raw_text).
        output_path: Absolute path for the output .xlsx file.

    Returns:
        The output_path on success.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Voters"

    # ── Header styling ───────────────────────────────────────
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write header row
    for col_idx, col_name in enumerate(FINAL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # ── Data rows ────────────────────────────────────────────
    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center", wrap_text=True)

    for row_idx, row_data in enumerate(all_results, 2):
        for col_idx, col_name in enumerate(FINAL_COLUMNS, 1):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=row_data.get(col_name, ""),
            )
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border

    # ── Auto-fit column widths ───────────────────────────────
    for col_idx, col_name in enumerate(FINAL_COLUMNS, 1):
        max_len = len(str(col_name))
        for row_idx in range(2, min(len(all_results) + 2, 102)):  # Sample first 100 rows
            val = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, len(val))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(
            max_len + 4, 50
        )

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(output_path)
    logger.info("Excel saved: %s (%d rows)", output_path, len(all_results))
    return output_path
